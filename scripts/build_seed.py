#!/usr/bin/env python3
"""Build the tracker seed from the two final program workbooks.

Sources of truth (2026-10 hand-over from the team's manual tracking):
  * CED TEMU US Referral Program.xlsx      -> program temu-us
  * CedCommerceUK_Agency ... -UKI.xlsx      -> program temu-uk
  * uk_sellers_import.sql (April 2026)      -> UK prospect baseline (campaign leads);
                                               every sheet value overrides it.

Conversion rules (confirmed by Krishna):
  * US: "breaking the first sale (Y or N)" == Y  -> eligible / converted
  * UK: row filled green (D9EAD3) in 'Seller status UK' -> converted

Output: data/seed/seed.json (+ seed_report.md). data/ is gitignored (PII).
Run:   python3 -I scripts/build_seed.py
"""
import datetime as dt
import json
import re
import sys
import uuid
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DL = Path.home() / 'Downloads'
US_XLSX = DL / 'CED TEMU US Referral Program.xlsx'
UK_XLSX = DL / 'CedCommerceUK_Agency and steve-tt-CedCommerceUK_agency-UKI.xlsx'
OLD_SQL = ROOT / 'uk_sellers_import.sql'
OUT = ROOT / 'data' / 'seed'

GREEN = 'FFD9EAD3'
NS = uuid.UUID('6f1c2a52-3c1e-4f43-9d0a-0d7c2b1e5a10')  # stable ids for sheet-only rows

ONBOARDED = 'ONBOARDING COMPLETE'
report = []


def log(msg):
    report.append(msg)


# ---------------------------------------------------------------- helpers
def txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r'[ \t]+', ' ', str(v).replace(' ', ' ')).strip()


def store_id(v):
    s = txt(v)
    m = re.search(r'mall_id=(\d+)', s)
    if m:
        return m.group(1)
    s = re.sub(r'\.0$', '', s)
    return s if re.fullmatch(r'\d{9,}', s) else ''


EMAIL_RE = re.compile(r'[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}')


def emails(v):
    """All emails in a messy cell, in order. 'wrong email - a / new email - b' -> new first."""
    s = txt(v)
    found = [e.lower() for e in EMAIL_RE.findall(s)]
    if 'new email' in s.lower() and len(found) > 1:
        found = found[1:] + found[:1]
    return list(dict.fromkeys(found))


def norm_name(v):
    return re.sub(r'[^a-z0-9]', '', txt(v).lower())


def sheet_date(v):
    """Messy first-sale dates: datetime, '9.09.25', '18-07-25', '13-12-25', '.'."""
    if isinstance(v, dt.datetime):
        return v.date()
    s = txt(v)
    m = re.fullmatch(r'(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})', s)
    if m:
        d, mo, y = map(int, m.groups())
        y = y + 2000 if y < 100 else y
        try:
            return dt.date(y, mo, d)
        except ValueError:
            return None
    return None


def clean_assoc(v):
    return txt(v).lstrip('@').strip()


def sid_uuid(*parts):
    return str(uuid.uuid5(NS, '|'.join(parts)))


# ---------------------------------------------------------------- old SQL baseline
def parse_old_sql(path):
    """Tokenise the Postgres INSERT tuples in uk_sellers_import.sql."""
    src = path.read_text()
    cols = ['id', 'program_id', 'email', 'company', 'status', 'sources', 'associate',
            'store_id', 'tags', 'phone', 'remarks', 'website', 'qualified']
    rows, i, n = [], 0, len(src)

    def read_value(j):
        while src[j] in ' \n\t':
            j += 1
        if src[j] == "'":
            j += 1
            buf = []
            while True:
                if src[j] == "'" and src[j + 1] == "'":
                    buf.append("'"); j += 2
                elif src[j] == "'":
                    return ''.join(buf), j + 1
                else:
                    buf.append(src[j]); j += 1
        if src.startswith('ARRAY[', j):
            j += 6
            items = []
            while src[j] != ']':
                if src[j] == "'":
                    v, j = read_value(j)
                    items.append(v)
                else:
                    j += 1
            j += 1
            if src.startswith('::TEXT[]', j):
                j += 8
            return items, j
        m = re.match(r'(NULL|true|false|-?\d+(?:\.\d+)?)', src[j:])
        tok = m.group(1)
        val = None if tok == 'NULL' else True if tok == 'true' else False if tok == 'false' else tok
        return val, j + len(tok)

    for block in re.finditer(r'INSERT INTO sellers \([^)]*\) VALUES', src):
        i = block.end()
        while True:
            while src[i] in ' \n\t':
                i += 1
            if src[i] != '(':
                break
            i += 1
            vals = []
            for _ in cols:
                v, i = read_value(i)
                vals.append(v)
                while src[i] in ' \n\t':
                    i += 1
                if src[i] == ',':
                    i += 1
            while src[i] != ')':
                i += 1
            i += 1
            rows.append(dict(zip(cols, vals)))
            while src[i] in ' \n\t':
                i += 1
            if src[i] == ',':
                i += 1
                continue
            break
    camps = []
    m = re.search(r'INSERT INTO campaigns \(([^)]*)\) VALUES(.*?)ON CONFLICT', src, re.S)
    if m:
        ccols = [c.strip() for c in m.group(1).split(',')]
        for t in re.findall(r"\((.*?)\)(?:,|\s*$)", m.group(2).strip(), re.S):
            parts = [p.strip() for p in re.findall(r"'(?:[^']|'')*'|[^,]+", t)]
            rec = {}
            for k, p in zip(ccols, parts):
                rec[k] = p[1:-1].replace("''", "'") if p.startswith("'") else int(p)
            camps.append(rec)
    return rows, camps


# ---------------------------------------------------------------- seller store
class Book:
    """Seller rows for one program, matched by store_id, then email, then name."""

    def __init__(self, program_id):
        self.pid = program_id
        self.rows = []

    def find(self, sid='', ems=(), name=''):
        for r in self.rows:
            if sid and r.get('store_id') == sid:
                return r
        for e in ems:
            for r in self.rows:
                if e and (r.get('email') == e or e in r['meta'].get('alt_emails', [])):
                    return r
        nn = norm_name(name)
        if nn and len(nn) > 3 and not re.fullmatch(r'[a-z]\d+shop', nn):
            for r in self.rows:
                if norm_name(r.get('company')) == nn:
                    return r
        return None

    def upsert(self, key, sid='', ems=(), name='', **fields):
        r = self.find(sid, ems, name)
        if r is None:
            r = {'id': sid_uuid(self.pid, key), 'program_id': self.pid, 'email': None, 'company': '',
                 'status': 'EMAIL FOLLOW UP IN PROGRESS', 'sources': [], 'associate': '', 'store_id': None,
                 'tags': [], 'phone': None, 'remarks': None, 'website': None, 'qualified': False, 'meta': {}}
            self.rows.append(r)
        if sid and not r.get('store_id'):
            r['store_id'] = sid
        for e in ems:
            if not r.get('email'):
                r['email'] = e
            elif e != r['email']:
                alts = r['meta'].setdefault('alt_emails', [])
                if e not in alts:
                    alts.append(e)
        if name and (not r.get('company') or re.fullmatch(r"[A-Z]\d+'shop", r['company'] or '')):
            r['company'] = name
        for k, v in fields.items():
            if v in (None, '', []):
                continue
            if k == 'sources':
                r['sources'] = list(dict.fromkeys(r['sources'] + v))
            elif k == 'remarks':
                r['remarks'] = v if not r.get('remarks') else (r['remarks'] if v in r['remarks'] else r['remarks'] + ' | ' + v)
            elif k == 'meta':
                r['meta'].update({mk: mv for mk, mv in v.items() if mv not in (None, '', [])})
            else:
                r[k] = v
        return r


def first_cell_fill(cells):
    for c in cells:
        f = c.fill
        if f and f.fill_type and f.fgColor is not None and f.fgColor.type == 'rgb':
            if f.fgColor.rgb == GREEN:
                return True
    return False


# ---------------------------------------------------------------- UK
def build_uk():
    book = Book('temu-uk')
    old, camps = parse_old_sql(OLD_SQL)
    old_status = Counter(r['status'] for r in old)
    for r in old:
        r = dict(r)
        r['meta'] = {'baseline': 'apr-2026 import'}
        r['sources'] = r['sources'] or []
        r['tags'] = r['tags'] or []
        r['email'] = (r['email'] or '').lower() or None
        r['store_id'] = store_id(r['store_id']) or None
        if r['status'] == 'ONBOARIDNG COMPLETED':
            r['status'] = ONBOARDED
        # Conversions are re-derived from the green rows below.
        if r['status'] == ONBOARDED:
            r['status'] = 'SALES AWAITING'
            r['meta']['was_onboarded_in_apr_import'] = True
        r['qualified'] = False
        book.rows.append(r)
    log(f'UK baseline: {len(old)} rows from uk_sellers_import.sql; statuses {dict(old_status)}')

    wb = openpyxl.load_workbook(UK_XLSX, data_only=True)

    # --- Seller status UK (Temu-confirmed stores; green = converted)
    ws = wb['Seller status UK']
    section = 'CedCommerceUK_agency'
    sale_dates = {}  # seller id -> every first-sale date the sheets give it (a store can appear on several rows)
    green_rows, seen_rows = [], 0
    for row in ws.iter_rows(min_row=4):
        a = txt(row[0].value)
        if a.startswith('steve-tt-CedCommerceUK'):
            section = 'steve-tt-CedCommerceUK_agency-UKI'
            continue
        sid = store_id(row[1].value) or store_id(row[2].value)
        if not sid:
            continue  # date / section banner rows
        seen_rows += 1
        name = a
        ems = emails(row[4].value) + emails(row[15].value)
        compliance = {k: txt(row[i].value) for k, i in [('entity', 7), ('tax', 8), ('shop', 9), ('identity', 10),
                                                         ('additional_compliance', 11), ('pre_withdrawal', 12)]}
        compliance = {k: v for k, v in compliance.items() if v}
        green = first_cell_fill(row[:19])
        fs = sheet_date(row[13].value)
        h = txt(row[7].value).lower()
        if green:
            status = ONBOARDED
        elif 'reject' in h or any('reject' in v.lower() for v in compliance.values()):
            status = 'Document Verification'
        elif any(w in h for w in ('approved', 'open', 'operation')):
            status = 'SALES AWAITING'
        else:
            status = 'REGISTRATION DONE'
        r = book.upsert(sid, sid=sid, ems=ems, name=name,
                        status=status, sources=['Marketplace Confirmed'],
                        remarks=txt(row[16].value),
                        meta={'temu_bd_manager': txt(row[6].value).split('\n')[0].strip(),
                              'vertical': txt(row[3].value),
                              'attention': txt(row[5].value).replace('duplicate entryswapnil.dhall@temu.com', 'duplicate entry (swapnil.dhall@temu.com)'),
                              'temu_store_status': txt(row[7].value),
                              'compliance': compliance or None,
                              'first_sale_date': fs.isoformat() if fs else None,
                              'incentivised_month': txt(row[14].value),
                              'ced_challenges': txt(row[17].value),
                              'agency_account': section,
                              'store_link': f'https://www.temu.com/mall.html?mall_id={sid}',
                              'sheet_row': row[0].row})
        r['status'] = status  # sheet is authoritative
        r['qualified'] = green
        if fs:
            sale_dates.setdefault(r['id'], []).append(fs)
        if green:
            green_rows.append((row[0].row, r))
    log(f'UK Seller status sheet: {seen_rows} store rows, {len(green_rows)} green (converted): '
        + ', '.join(f"{r['company']} (row {n})" for n, r in green_rows))

    # --- BD Tracking (associate attribution + funnel status for BD-sourced leads)
    ws = wb['BD Tracking ']
    bd_n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, '') for v in row[:6]):
            continue
        bd_n += 1
        ems = emails(row[1])
        site = txt(row[2])
        st = txt(row[3])
        st = {'Onboarding Completed': 'SALES AWAITING'}.get(st, st)  # conversion only via green
        r = book.upsert('bd:' + (ems[0] if ems else site), ems=ems, name='' if '.' in site else site,
                        associate=clean_assoc(row[0]), website=site if '.' in site else None,
                        sources=['BD referred'], remarks=' | '.join(x for x in (txt(row[4]), txt(row[5])) if x))
        if not r['qualified']:
            r['status'] = st
    log(f'UK BD Tracking: {bd_n} rows merged')

    # --- CedCommerce Sellers tab (Temu-confirmed list kept by CS)
    ws = wb['CedCommerce Sellers']
    for row in ws.iter_rows(min_row=2, values_only=True):
        sid = store_id(row[1])
        if not sid:
            continue
        r = book.upsert(sid, sid=sid, ems=emails(row[2]), name=txt(row[0]), sources=['Marketplace Confirmed'])
        if txt(row[3]) == 'Not In Touch' and 'Not in touch' not in r['tags']:
            r['tags'].append('No reply')

    # --- Paid tracking -> qualified_sellers (exactly the green rows)
    ws = wb['Ced Paid CSA']
    paid = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        nm = txt(row[1])
        if not nm:
            continue
        paid[norm_name(nm)] = {'month': row[0], 'store': nm, 'bda': clean_assoc(row[2]),
                               'ced_paid': txt(row[3]).lower(), 'csa_paid': txt(row[4]).lower(),
                               'note_month': txt(row[5])}
    inc_month = {}
    ws = wb['Incentives tracker']
    cur = None
    months = {'JULY': '2025-07', 'AUG': '2025-08', 'SEP:': '2025-09', 'OCT:': '2025-10', 'November': '2025-11',
              'January': '2026-01', 'February': '2026-02'}
    for row in ws.iter_rows(min_row=1, values_only=True):
        a = txt(row[0])
        if a in months:
            cur = months[a]
            continue
        if cur and (a or row[1]):
            for sid in re.findall(r'\d{12,}', re.sub(r'\.0\b', '', a)):
                inc_month[sid] = cur

    quals = []
    for _, r in green_rows:
        p = paid.get(norm_name(r['company'])) or next(
            (v for k, v in paid.items() if norm_name(r['company'])[:8] and k.startswith(norm_name(r['company'])[:8])), None)
        if p is None and 'supremecbd' in norm_name(r['company']):
            p = paid.get(norm_name("A245334'shop(supremecbd-shop.myshopify.com)"))
        # Rule (Krishna, 2026-10-12): when the sheets give several dates, take the latest one.
        # Dates after today can't be a first sale yet, so they're ignored (and noted).
        cands = list(sale_dates.get(r['id'], []))
        if p and isinstance(p['month'], dt.datetime):
            cands.append(p['month'].date())
        today = dt.date.today()
        future = sorted(d for d in cands if d > today)
        valid = sorted(d for d in cands if d <= today)
        fs = valid[-1].isoformat() if valid else None
        fsm = fs[:7] if fs else None
        r['meta']['first_sale_date'] = fs
        if p and p['bda']:
            r['associate'] = p['bda']
        notes = []
        if r['store_id'] in inc_month:
            notes.append(f"Incentive claim month {inc_month[r['store_id']]}")
        if len(cands) > 1:
            notes.append('Sheet dates: ' + ', '.join(d.isoformat() for d in sorted(cands)) + f' → latest used {fs}')
        if future:
            notes.append('Ignored future date(s): ' + ', '.join(d.isoformat() for d in future))
        quals.append({'id': sid_uuid('qual', r['program_id'], r['store_id'] or r['company']),
                      'program_id': 'temu-uk', 'store_name': r['company'], 'email': r['email'],
                      'store_id': r['store_id'], 'bda': (p or {}).get('bda') or r['associate'] or None,
                      'first_sale_month': fsm, 'amount': 800, 'currency': 'GBP',
                      'temu_paid_ced': 'paid' if p and p['ced_paid'] == 'paid' else 'pending',
                      'ced_paid_csa': 'paid' if p and p['csa_paid'] == 'paid' else 'pending',
                      'bd_incentive': 'pending', 'notes': '; '.join(notes) or None})
        if not p:
            log(f"WARN UK green row {r['company']} has no 'Ced Paid CSA' entry")

    # --- payments: one row per incentive claim month (amount = £800 x stores Temu has paid)
    pays = []
    by_month = {}
    for q in quals:
        m = inc_month.get(q['store_id'])
        if m:
            by_month.setdefault(m, []).append(q)
    for m, qs in sorted(by_month.items()):
        paid_qs = [q for q in qs if q['temu_paid_ced'] == 'paid']
        pend = [q['store_name'] for q in qs if q['temu_paid_ced'] != 'paid']
        pays.append({'id': sid_uuid('pay', 'temu-uk', m), 'program_id': 'temu-uk', 'period': m,
                     'amount': 800 * len(paid_qs), 'seller_count': len(paid_qs), 'received': bool(paid_qs),
                     'stores_note': ' · '.join(x for x in (', '.join(q['store_name'] for q in paid_qs),
                                                           'pending: ' + ', '.join(pend) if pend else '') if x)})
    unclaimed = [q['store_name'] for q in quals if q['store_id'] not in inc_month]
    if unclaimed:
        log('UK converted but not yet in Incentives tracker: ' + ', '.join(unclaimed))

    campaigns = [{**c, 'program_id': 'temu-uk'} for c in camps]
    return book.rows, quals, pays, campaigns


# ---------------------------------------------------------------- US
def us_status(store_status, bd_status, y):
    s = store_status.lower()
    if y:
        return ONBOARDED
    if 'will not' in s or 'no action is needed' in s:
        return 'NOT INTERESTED'
    if bd_status and bd_status not in ('Onboarding Completed',):
        return bd_status
    if s.startswith('live') or 'live,' in s:
        return 'SALES AWAITING'
    if 'error' in s or 'missing' in s:
        return 'Document Verification'
    if 'no submission' in s:
        return 'REGISTRATION DONE'
    return 'REGISTRATION DONE'


def build_us():
    book = Book('temu-us')
    wb = openpyxl.load_workbook(US_XLSX, data_only=True)
    info = {txt(r[0]).rstrip(' -').strip(): txt(r[1]) for r in wb['Program Info'].iter_rows(values_only=True)}
    benefits = [txt(r[1]) for r in wb['Program Info'].iter_rows(min_row=3, values_only=True) if txt(r[1])]

    bd = []
    for row in wb['BD Tracking '].iter_rows(min_row=2, values_only=True):
        if any(v not in (None, '') for v in row[:6]):
            bd.append({'assoc': clean_assoc(row[0]), 'ems': emails(row[1]), 'site': txt(row[2]),
                       'status': txt(row[3]), 'remarks': txt(row[4]), 'temu_feedback': txt(row[5]),
                       'sid': store_id(row[6])})
    forms = {}
    for row in wb['Form Responses (US)'].iter_rows(min_row=2, values_only=True):
        ems = emails(row[0])
        if ems:
            forms[ems[0]] = {'background': txt(row[5]), 'issue': txt(row[2]), 'end_to_end': txt(row[6]),
                             'submitted_by': txt(row[7]), 'submitted_at': txt(row[8])}

    quals, y_n = [], 0
    ws = wb['Temu Seller Status US']
    for row in ws.iter_rows(min_row=2, values_only=True):
        sid = store_id(row[3])
        if not sid:
            continue
        ems = emails(row[1])
        b = next((x for x in bd if x['sid'] == sid or set(x['ems']) & set(ems)), None)
        if b:
            ems = ems + [e for e in b['ems'] if e not in ems]
        y = txt(row[8]).upper() == 'Y'
        y_n += y
        name = txt(row[4])
        status = us_status(txt(row[9]), b['status'] if b else '', y)
        phone = txt(row[12]) if re.search(r'\d{3}.*\d{4}', txt(row[12])) else None
        r = book.upsert(sid, sid=sid, ems=ems, name=name, status=status,
                        sources=['Marketplace Confirmed'] + (['BD referred'] if b else []),
                        associate=b['assoc'] if b else '', phone=phone,
                        website=(b['site'] if b and '.' in b['site'] else None),
                        remarks=b['remarks'] if b else '',
                        meta={'contact_name': txt(row[0]), 'payout': txt(row[2]), 'entity_id': store_id(row[5]),
                              'listings': txt(row[6]), 'live_listings': txt(row[7]),
                              'first_sale': 'Y' if y else txt(row[8]) or None,
                              'temu_store_status': txt(row[9]), 'temu_am_comments': txt(row[10]),
                              'whatsapp_with_temu_am': txt(row[11]),
                              'temu_feedback': b['temu_feedback'] if b else '',
                              'legal_name': b['site'] if b and '.' not in b['site'] else '',
                              'store_link': f'https://www.temu.com/mall.html?mall_id={sid}'})
        r['status'] = status
        r['qualified'] = y
        if y:
            payout = txt(row[2])
            m = re.search(r'(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)', payout, re.I)
            pm = f"2026-{dt.datetime.strptime(m.group(1)[:3].title(), '%b').month:02d}" if m else None
            quals.append({'id': sid_uuid('qual', 'temu-us', sid), 'program_id': 'temu-us', 'store_name': name,
                          'email': r['email'], 'store_id': sid, 'bda': r['associate'] or None,
                          'first_sale_month': None, 'amount': 1000, 'currency': 'USD',
                          'temu_paid_ced': 'paid' if payout.upper().startswith('Y') else 'pending',
                          'ced_paid_csa': 'pending', 'bd_incentive': 'pending',
                          'notes': f'Temu payout: {payout}' + (f' ({pm})' if pm else '')})
    # BD-only leads (not yet visible to Temu)
    for b in bd:
        if any(set(b['ems']) & ({r['email']} | set(r['meta'].get('alt_emails', []))) for r in book.rows):
            continue
        st = b['status'] if b['status'] != 'Onboarding Completed' else 'SALES AWAITING'
        book.upsert('bd:' + (b['ems'][0] if b['ems'] else b['site']), ems=b['ems'],
                    name=b['site'] if '.' not in b['site'] else '', status=st, associate=b['assoc'],
                    website=b['site'] if '.' in b['site'] else None, sources=['BD referred'],
                    remarks=' | '.join(x for x in (b['remarks'],) if x),
                    meta={'temu_feedback': b['temu_feedback']})
    for r in book.rows:
        f = forms.get(r['email'] or '')
        if f:
            r['meta']['bd_form'] = f
    log(f"US: {len(book.rows)} sellers ({sum(1 for r in book.rows if 'Marketplace Confirmed' in r['sources'])} "
        f"Temu-confirmed), {y_n} first-sale = Y: " + ', '.join(q['store_name'] for q in quals))
    program = {'id': 'temu-us', 'name': 'Temu US', 'market': 'Temu', 'region': 'US', 'incentive_amount': 1000,
               'currency': 'USD', 'color': 'blue', 'affiliate_link': info.get('Affiliate Link') or None,
               'notes': 'Start date: ' + info.get('Start Date', '') + '\nSeller benefits:\n- ' + '\n- '.join(benefits),
               'meta': {'start_date': '2026-02-16', 'seller_benefits': benefits,
                        'temu_contact': 'Candice Dai (candice.dai@temu.com)',
                        'eligibility_rule': 'First sale = Y in Temu Seller Status US'}}
    return program, book.rows, quals


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    uk_rows, uk_q, uk_p, uk_c = build_uk()
    us_prog, us_rows, us_q = build_us()
    uk_prog = {'id': 'temu-uk', 'name': 'Temu UK', 'market': 'Temu', 'region': 'UK', 'incentive_amount': 800,
               'currency': 'GBP', 'color': 'teal', 'affiliate_link': None,
               'notes': "Temu UK agency/consultant incentive: £800 per new seller that registers via CedCommerce's "
                        'agency link and makes at least one sale outside a test order. Temu pays on the 15th; '
                        'payment request goes on the 3rd of each month.',
               'meta': {'agency_accounts': ['CedCommerceUK_agency', 'steve-tt-CedCommerceUK_agency-UKI'],
                        'eligibility_rule': "Green row in 'Seller status UK'"}}
    shein = ROOT / 'data' / 'shein-ab-brands'
    shein_prog = json.loads((shein / 'program.json').read_text())
    shein_sellers = json.loads((shein / 'sellers_seed.json').read_text())
    seed = {'programs': [uk_prog, us_prog],
            'sellers': uk_rows + us_rows,
            'qualified_sellers': uk_q + us_q,
            'payments': uk_p,
            'campaigns': uk_c,
            '_shein': {'program': shein_prog, 'sellers': shein_sellers}}
    # sanity
    ids = Counter(r['id'] for r in seed['sellers'])
    dupes = [k for k, v in ids.items() if v > 1]
    assert not dupes, f'duplicate ids: {dupes[:5]}'
    for p in ('temu-uk', 'temu-us'):
        rows = [r for r in seed['sellers'] if r['program_id'] == p]
        st = Counter(r['status'] for r in rows)
        log(f'{p}: {len(rows)} sellers, {sum(r["qualified"] for r in rows)} converted; statuses {dict(st.most_common())}')
    (OUT / 'seed.json').write_text(json.dumps(seed, indent=1, default=str))
    (OUT / 'seed_report.md').write_text('\n'.join('- ' + m for m in report) + '\n')
    print('\n'.join(report))


if __name__ == '__main__':
    sys.exit(main())
